import os
import re
import io
import pandas as pd
import requests
import streamlit as st
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------
# STREAMLIT PAGE CONFIG
# ---------------------------------------------------------
st.set_page_config(
    page_title="Journal Metrics & Quartile Lookup",
    page_icon="📚",
    layout="wide"
)

st.title("📚 Journal Metrics & Quartile Lookup Tool")
st.write("Lookup SCImago Quartiles (Q1–Q4), SJR Scores, and 2-Year Impact Factor equivalents by Journal Name, ISSN, or reference files (`.ris`, `.bib`, `.xlsx`, `.csv`).")

# ---------------------------------------------------------
# CACHED DATA LOADERS & HELPERS
# ---------------------------------------------------------
SCIMAGO_FILE = "scimago_data.csv"

def parse_issns(issn_str):
    if pd.isna(issn_str):
        return []
    return [i.strip() for i in str(issn_str).replace("-", "").replace(",", " ").split() if i.strip()]

def standardize_scimago_columns(df):
    if df is None:
        return None
    col_mapping = {col.strip().lower(): col for col in df.columns}

    title_key = next((col_mapping[k] for k in ['title', 'journal title'] if k in col_mapping), None)
    q_key = next((col_mapping[k] for k in ['sjr best q', 'sjr best quartile', 'quartile'] if k in col_mapping), None)
    issn_key = next((col_mapping[k] for k in ['issn'] if k in col_mapping), None)
    sjr_key = next((col_mapping[k] for k in ['sjr'] if k in col_mapping), None)
    if_key = next((col_mapping[k] for k in ['cites / doc. (2years)', 'cites/doc. (2years)'] if k in col_mapping), None)

    if title_key and q_key:
        rename_dict = {title_key: 'Title', q_key: 'SJR Best Q'}
        if issn_key: rename_dict[issn_key] = 'ISSN'
        if sjr_key: rename_dict[sjr_key] = 'SJR Score'
        if if_key: rename_dict[if_key] = 'Impact Factor Equivalent (2y)'
        return df.rename(columns=rename_dict)
    return None

@st.cache_data(show_spinner="Loading SCImago Dataset...")
def load_scimago_database():
    if not os.path.exists(SCIMAGO_FILE):
        try:
            scimago_url = "https://www.scimagojr.com/journalrank.php?out=xls"
            headers = {
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/115.0.0.0 Safari/537.36'
            }
            response = requests.get(scimago_url, headers=headers, timeout=15)
            response.raise_for_status()
            with open(SCIMAGO_FILE, 'wb') as f:
                f.write(response.content)
        except Exception as e:
            st.error(f"Failed to download SCImago data automatically: {e}")
            return None, {}, {}

    try:
        df_scimago = pd.read_csv(SCIMAGO_FILE, sep=";")
    except Exception:
        try:
            df_scimago = pd.read_csv(SCIMAGO_FILE, sep=";", encoding="latin-1")
        except Exception:
            return None, {}, {}

    df_scimago = standardize_scimago_columns(df_scimago)
    if df_scimago is None:
        return None, {}, {}

    df_scimago["Clean_Title"] = df_scimago["Title"].astype(str).str.strip().str.lower()

    issn_metrics_map = {}
    journal_metrics_map = {}

    for idx, row in df_scimago.iterrows():
        quartile = str(row["SJR Best Q"]) if pd.notna(row["SJR Best Q"]) else ""
        sjr_score = row.get("SJR Score", "")
        impact_factor = row.get("Impact Factor Equivalent (2y)", "")

        try:
            if pd.notna(sjr_score): sjr_score = float(str(sjr_score).replace(",", "."))
        except ValueError: pass
        try:
            if pd.notna(impact_factor): impact_factor = float(str(impact_factor).replace(",", "."))
        except ValueError: pass

        metrics = {"Quartile": quartile, "SJR": sjr_score, "ImpactFactor_2y": impact_factor}

        journal_metrics_map[row["Clean_Title"]] = metrics
        if "ISSN" in df_scimago.columns:
            for issn in parse_issns(row["ISSN"]):
                issn_metrics_map[issn] = metrics

    return df_scimago, journal_metrics_map, issn_metrics_map

def lookup_metrics(query, journal_map, issn_map):
    default_metrics = {"Quartile": "", "SJR": "", "ImpactFactor_2y": ""}
    if not query or pd.isna(query):
        return default_metrics

    clean_query = str(query).strip().lower().replace("-", "")
    if clean_query.isdigit() and len(clean_query) == 8:
        if clean_query in issn_map:
            return issn_map[clean_query]

    clean_title_query = str(query).strip().lower()
    if clean_title_query in journal_map:
        return journal_map[clean_title_query]

    return default_metrics

# ---------------------------------------------------------
# FILE PARSERS
# ---------------------------------------------------------
def parse_ris_bytes(file_bytes):
    try: text = file_bytes.decode('utf-8')
    except UnicodeDecodeError: text = file_bytes.decode('latin-1')
    
    entries, current_entry = [], {}
    for line in text.splitlines():
        line = line.strip()
        if not line: continue
        match = re.match(r'^([A-Z0-9]{2})\s*-\s*(.*)$', line)
        if match:
            tag, value = match.groups()
            if tag == 'TY':
                if current_entry: entries.append(current_entry)
                current_entry = {'Type': value}
            elif tag == 'ER':
                if current_entry:
                    entries.append(current_entry)
                    current_entry = {}
            else:
                if current_entry is not None:
                    if tag in current_entry:
                        if isinstance(current_entry[tag], list): current_entry[tag].append(value)
                        else: current_entry[tag] = [current_entry[tag], value]
                    else: current_entry[tag] = value
    if current_entry: entries.append(current_entry)

    structured = []
    for entry in entries:
        title = entry.get('T1', entry.get('TI', ''))
        if isinstance(title, list): title = " ".join(title)
        year = entry.get('PY', entry.get('Y1', ''))
        if isinstance(year, list): year = year[0]
        ym = re.search(r'\b\d{4}\b', str(year))
        year = ym.group(0) if ym else year
        journal = entry.get('JO', entry.get('JF', entry.get('JA', entry.get('T2', ''))))
        if isinstance(journal, list): journal = journal[0]
        authors = entry.get('AU', '')
        if isinstance(authors, list): authors = "; ".join(authors)
        issn = entry.get('SN', '')
        if isinstance(issn, list): issn = issn[0]
        doi = entry.get('DO', entry.get('M3', ''))
        if isinstance(doi, list): doi = doi[0]
        publisher = entry.get('PB', '')
        if isinstance(publisher, list): publisher = "; ".join(publisher)

        structured.append({
            'Title': title, 'Author': authors, 'Year': year,
            'JournalName': journal, 'ISSN': issn, 'DOI': doi, 'Publisher': publisher
        })
    return pd.DataFrame(structured)

def parse_bib_bytes(file_bytes):
    try: content = file_bytes.decode('utf-8')
    except UnicodeDecodeError: content = file_bytes.decode('latin-1')

    raw_entries = re.findall(r'@\w+\s*\{\s*([^,]+),([\s\S]*?)\n\s*\}(?=\s*@|\s*$)', content)
    structured = []
    for key, body in raw_entries:
        fields = re.findall(r'(\w+)\s*=\s*([{"‘’“”].*?[}"‘’“”]|[\w\d-]+)', body, re.DOTALL)
        entry_dict = {}
        for field, value in fields:
            clean_val = value.strip().strip('{}\u201c\u201d\u2018\u2019"\'')
            entry_dict[field.strip().lower()] = re.sub(r'\s+', ' ', clean_val)

        structured.append({
            'Title': entry_dict.get('title', ''),
            'Author': entry_dict.get('author', '').replace(' and ', '; '),
            'Year': entry_dict.get('year', ''),
            'JournalName': entry_dict.get('journal', entry_dict.get('booktitle', '')),
            'ISSN': entry_dict.get('issn', ''),
            'DOI': entry_dict.get('doi', ''),
            'Publisher': entry_dict.get('publisher', '')
        })
    return pd.DataFrame(structured)

def generate_styled_excel(df):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name="Journals")
        wb = writer.book
        ws = writer.sheets["Journals"]
        
        max_row, max_col = ws.max_row, ws.max_column
        ref_range = f"A1:{get_column_letter(max_col)}{max_row}"
        
        tab = Table(displayName="MetricsTable", ref=ref_range)
        tab.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False
        )
        ws.add_table(tab)
    return output.getvalue()

# ---------------------------------------------------------
# MAIN APP ROUTING
# ---------------------------------------------------------
df_scimago, journal_map, issn_map = load_scimago_database()

if df_scimago is None:
    st.error("SCImago database could not be loaded. Please check your internet connection or upload the file manually.")
else:
    st.sidebar.header("Navigation")
    option = st.sidebar.radio("Choose Mode:", ["Single Journal Search", "Batch Citation Lookup"])

    if option == "Single Journal Search":
        st.subheader("🔎 Single Journal / ISSN Lookup")
        query_input = st.text_input("Enter Journal Title or ISSN (e.g., '0040-5175' or 'Textile Research Journal'):")
        
        if st.button("Search Metrics", type="primary"):
            if query_input:
                metrics = lookup_metrics(query_input, journal_map, issn_map)
                if metrics["Quartile"] or metrics["SJR"]:
                    st.success("Record Found!")
                    col1, col2, col3 = st.columns(3)
                    col1.metric("Quartile", metrics["Quartile"] if metrics["Quartile"] else "N/A")
                    col2.metric("SJR Score", metrics["SJR"] if metrics["SJR"] else "N/A")
                    col3.metric("2-Year IF Equivalent", metrics["ImpactFactor_2y"] if metrics["ImpactFactor_2y"] else "N/A")
                else:
                    st.warning("Journal or ISSN not found in the SCImago dataset.")
            else:
                st.info("Please enter a journal name or ISSN to search.")

    elif option == "Batch Citation Lookup":
        st.subheader("📁 Batch File Processing")
        uploaded_file = st.file_uploader("Upload a file (`.ris`, `.bib`, `.xlsx`, `.csv`)", type=["ris", "bib", "xlsx", "csv"])
        
        if uploaded_file is not None:
            filename = uploaded_file.name
            df_user = pd.DataFrame()

            if filename.endswith(".ris"):
                df_user = parse_ris_bytes(uploaded_file.getvalue())
            elif filename.endswith(".bib"):
                df_user = parse_bib_bytes(uploaded_file.getvalue())
            elif filename.endswith(".csv"):
                df_user = pd.read_csv(uploaded_file)
            elif filename.endswith(".xlsx"):
                df_user = pd.read_excel(uploaded_file)

            if not df_user.empty:
                st.write(f"**Loaded {len(df_user)} records from `{filename}`**")
                
                # Column selector if standard JournalName column isn't present
                target_col = "JournalName"
                if target_col not in df_user.columns:
                    target_col = st.selectbox("Select the column containing Journal Names or ISSNs:", df_user.columns)

                if st.button("Process & Resolve Metrics", type="primary"):
                    def fetch_metrics(row):
                        if 'ISSN' in row and pd.notna(row['ISSN']) and str(row['ISSN']).strip():
                            res = lookup_metrics(row['ISSN'], journal_map, issn_map)
                            if res["Quartile"]: return res
                        return lookup_metrics(row[target_col], journal_map, issn_map)

                    resolved = df_user.apply(fetch_metrics, axis=1)
                    df_user["Quartile"] = [m["Quartile"] for m in resolved]
                    df_user["SJR Score"] = [m["SJR"] for m in resolved]
                    df_user["Impact Factor Equivalent (2y)"] = [m["ImpactFactor_2y"] for m in resolved]

                    st.dataframe(df_user, use_container_width=True)

                    excel_data = generate_styled_excel(df_user)
                    st.download_button(
                        label="📥 Download Excel File with Metrics",
                        data=excel_data,
                        file_name="Updated_Journals_with_Quartiles.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
            else:
                st.error("Could not extract valid records from the uploaded file.")